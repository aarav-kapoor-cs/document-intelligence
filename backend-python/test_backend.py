"""Unit tests for the Excel audit export and the database helper.

Run from this folder with:  python -m unittest test_backend -v
They use SQLite in a temporary file and never touch Azure or SQL Server.
"""
import io
import os
import tempfile
import unittest
from datetime import date, datetime

import openpyxl

import database
import excel_service

GSTIN = "27ABICX1218R1ZX"
PAN = "ABICX1218R"


def record(id, file_name, fields, created_at="2026-07-23 10:00:00"):
    """A record shaped exactly like database.get_filtered() returns."""
    return {
        "id": id, "fileName": file_name, "model": "prebuilt-invoice",
        "fields": fields, "aiAnswerJson": None, "createdAt": created_at,
    }


def wrap(value):
    """A field exactly like ocr_service stores it in KeyValuesJson."""
    return {"value": value, "confidence": 0.9}


class HelperTests(unittest.TestCase):
    def test_is_blank(self):
        for value in (None, "", "  ", "none", "null", [], {}, "[]", "{}"):
            self.assertTrue(excel_service.is_blank(value), repr(value))
        for value in ("0", 0, "x", ["a"], {"a": 1}):
            self.assertFalse(excel_service.is_blank(value), repr(value))

    def test_plain_unwraps_only_ocr_wrappers(self):
        self.assertEqual(excel_service._plain(wrap("46273")), "46273")
        # A real business dict that merely contains "value" must stay intact.
        data = {"value": 100, "currency": "INR"}
        self.assertEqual(excel_service._plain(data), data)

    def test_gstin_and_pan_rules(self):
        self.assertTrue(excel_service.is_valid_gstin(GSTIN))
        self.assertTrue(excel_service.is_valid_gstin(" 27abicx1218r1zx "))  # normalized
        self.assertFalse(excel_service.is_valid_gstin(PAN))
        self.assertTrue(excel_service.is_pan(PAN))

    def test_vendor_gstin_priority_and_pan_rejection(self):
        # GSTIN inside a lower-priority key still wins over a PAN in a higher one.
        fields = {"VendorTaxId": wrap(PAN), "GSTIN/UIN": wrap(GSTIN)}
        self.assertEqual(excel_service.get_vendor_gstin_from_key_values_json(fields), GSTIN)
        # A PAN alone must never come back.
        self.assertEqual(
            excel_service.get_vendor_gstin_from_key_values_json({"VendorTaxId": wrap(PAN)}), "")

    def test_parse_date_formats(self):
        for text in ("2025-04-16", "16/04/2025", "16-04-2025", "16 Apr 2025"):
            self.assertEqual(excel_service.parse_date(text), date(2025, 4, 16), text)
        self.assertIsNone(excel_service.parse_date("not a date"))

    def test_extraction_date_is_local_ist(self):
        # 19:30 UTC = 01:00 IST the NEXT day.
        self.assertEqual(
            excel_service.extraction_date_local("2026-07-23 19:30:00"), date(2026, 7, 24))
        self.assertEqual(
            excel_service.extraction_date_local("2026-07-23 10:00:00"), date(2026, 7, 23))
        self.assertIsNone(excel_service.extraction_date_local("garbage"))

    def test_is_valid_field_value(self):
        self.assertTrue(excel_service.is_valid_field_value("VendorTaxId", GSTIN))
        self.assertFalse(excel_service.is_valid_field_value("VendorTaxId", PAN))
        self.assertTrue(excel_service.is_valid_field_value("InvoiceDate", "16/04/2025"))
        self.assertFalse(excel_service.is_valid_field_value("InvoiceDate", "soon"))
        self.assertTrue(excel_service.is_valid_field_value("InvoiceTotal", "1,234.50"))
        self.assertFalse(excel_service.is_valid_field_value("InvoiceTotal", "n/a"))
        self.assertFalse(excel_service.is_valid_field_value("Anything", ""))


class ExtractionLogTests(unittest.TestCase):
    def setUp(self):
        self.records = [
            record(1, "a.pdf", {"InvoiceId": wrap("INV-1"), "VendorTaxId": wrap(GSTIN)}),
            record(2, "b.pdf", {"VendorTaxId": wrap(PAN)}),
        ]
        self.rows = excel_service.build_extraction_log_rows(self.records, "prebuilt-invoice")

    def row(self, file_name, field_name):
        return next(r for r in self.rows
                    if r["source_file"] == file_name and r["field_name"] == field_name)

    def test_one_row_per_union_field_per_record(self):
        # Union of fields = {InvoiceId, VendorTaxId} -> 2 rows for each record.
        self.assertEqual(len(self.rows), 4)

    def test_missing_field_is_incomplete(self):
        missing = self.row("b.pdf", "InvoiceId")
        self.assertEqual(missing["completeness_flag"], "N")
        self.assertEqual(missing["validity_flag"], "Fail")

    def test_pan_stays_visible_but_fails_validity(self):
        pan_row = self.row("b.pdf", "VendorTaxId")
        self.assertEqual(pan_row["ocr_value"], PAN)          # not blanked
        self.assertEqual(pan_row["completeness_flag"], "Y")  # a value exists
        self.assertEqual(pan_row["validity_flag"], "Fail")   # but it is not a GSTIN

    def test_valid_gstin_passes(self):
        good = self.row("a.pdf", "VendorTaxId")
        self.assertEqual(good["ocr_value"], GSTIN)
        self.assertEqual(good["validity_flag"], "Pass")

    def test_user_edit_always_blank(self):
        self.assertTrue(all(r["user_edit"] == "" for r in self.rows))


class TrendAndMixTests(unittest.TestCase):
    def test_percentages_are_decimals(self):
        records = [
            record(1, "a.pdf", {"InvoiceId": wrap("INV-1")}),
            record(2, "b.pdf", {"InvoiceId": wrap("INV-2"), "SubTotal": wrap("100")}),
        ]
        rows = excel_service.build_extraction_log_rows(records, "prebuilt-invoice")
        completeness = {c["ocr_field"]: c for c in excel_service.calculate_completeness_trend(rows)}
        self.assertEqual(completeness["Base Amount"]["completeness_percent"], 0.5)
        self.assertEqual(completeness["Invoice Id"]["completeness_percent"], 1.0)
        validity = {v["field_type"]: v for v in excel_service.calculate_validity_trend(rows)}
        self.assertEqual(validity["Base Amount"]["validity_percent"], 0.5)

    def test_invoice_volume_groups_by_ist_day_and_skips_garbage(self):
        records = [
            record(1, "a.pdf", {}, created_at="2026-07-23 19:30:00"),  # 24/07 in IST
            record(2, "b.pdf", {}, created_at="2026-07-23 10:00:00"),
            record(3, "c.pdf", {}, created_at="garbage"),
        ]
        volume = excel_service.calculate_invoice_volume(records)
        self.assertEqual(
            [(v["extraction_date"], v["invoice_count"]) for v in volume],
            [(date(2026, 7, 23), 1), (date(2026, 7, 24), 1)])

    def test_vendor_mix_decimal_and_never_pan(self):
        records = [
            record(1, "a.pdf", {"VendorTaxId": wrap(GSTIN)}),
            record(2, "b.pdf", {"VendorTaxId": wrap(GSTIN)}),
            record(3, "c.pdf", {"VendorTaxId": wrap(PAN)}),
            record(4, "d.pdf", {}),
        ]
        mix = {m["vendor_gstin"]: m for m in excel_service.calculate_vendor_mix(records)}
        self.assertEqual(mix[GSTIN]["invoice_count"], 2)
        self.assertEqual(mix[GSTIN]["vendor_percentage"], 0.5)
        self.assertEqual(mix["Unknown"]["invoice_count"], 2)  # PAN-only + empty
        self.assertNotIn(PAN, mix)


class WorkbookTests(unittest.TestCase):
    def build(self, records):
        content = excel_service.build_workbook(records, "2026-07-20", "2026-07-24", "prebuilt-invoice")
        return openpyxl.load_workbook(io.BytesIO(content))

    def test_six_sheets_with_verbatim_headers(self):
        wb = self.build([record(1, "a.pdf", {"InvoiceId": wrap("INV-1")})])
        self.assertEqual(wb.sheetnames, [
            "AUDIT Period", "AI EXTRACTION LOG", "COMPLETENESS TREND",
            "VALIDITY TREND", "Invoice Volume", "VENDOR MIX"])
        log_headers = [c.value for c in wb["AI EXTRACTION LOG"][1]]
        self.assertEqual(log_headers[3], "Field_Name(Keys from json) ")  # template quirk
        self.assertEqual([c.value for c in wb["VALIDITY TREND"][1]][1], "Type ")  # template quirk
        self.assertEqual(wb["Invoice Volume"]["A1"].value, "Extraction_Date")

    def test_dates_and_percents_formatted(self):
        wb = self.build([record(1, "a.pdf", {"InvoiceId": wrap("INV-1")})])
        audit = wb["AUDIT Period"]
        for cell in (audit["B2"], audit["C2"], audit["D2"]):
            self.assertTrue(cell.is_date)
            self.assertEqual(cell.number_format, "DD/MM/YYYY")
        self.assertTrue(str(audit["G2"].value).startswith("Azure OpenAI"))
        pct = wb["COMPLETENESS TREND"].cell(2, 6)
        self.assertEqual(pct.number_format, "0.00%")
        self.assertLessEqual(pct.value, 1.0)

    def test_empty_export_still_builds_all_sheets(self):
        wb = self.build([])
        self.assertEqual(len(wb.sheetnames), 6)
        self.assertEqual(wb["AUDIT Period"]["E2"].value, 0)


class DatabaseTests(unittest.TestCase):
    def setUp(self):
        # Point the module at a throwaway SQLite file.
        self.tmp = tempfile.NamedTemporaryFile(suffix=".db", delete=False)
        self.tmp.close()
        self.original_path = database.DB_PATH
        database.DB_PATH = self.tmp.name
        os.environ.pop("USE_SQL_SERVER", None)
        database.ensure_table()

    def tearDown(self):
        database.DB_PATH = self.original_path
        os.unlink(self.tmp.name)

    def test_window_utc_shifts_ist_to_utc(self):
        from_dt, to_dt = database._window_utc("2026-07-24", "2026-07-24")
        self.assertEqual(from_dt, datetime(2026, 7, 23, 18, 30, 0))
        self.assertEqual(to_dt, datetime(2026, 7, 24, 18, 29, 59))

    def test_save_ai_answer_json(self):
        self.assertIsNone(database.save_ai_answer_json(None))
        self.assertEqual(database.save_ai_answer_json({"a": 1}), '{"a": 1}')

    def test_save_and_get_filtered_round_trip(self):
        database.save("t.pdf", "prebuilt-invoice", "p", "text", '{"InvoiceId": "X"}',
                      '{"vendor": "ACME"}', "ans", 1, 2, 3, "2026-07-23 19:30:00", "pid")
        # 19:30 UTC on the 23rd is the 24th in IST -> found by the 24th's window.
        rows = database.get_filtered("2026-07-24", "2026-07-24", "prebuilt-invoice")
        self.assertEqual(len(rows), 1)
        self.assertEqual(rows[0]["fields"], {"InvoiceId": "X"})
        self.assertEqual(rows[0]["aiAnswerJson"], {"vendor": "ACME"})
        # ... and correctly NOT found by the 23rd-only IST window.
        self.assertEqual(database.get_filtered("2026-07-23", "2026-07-23", "prebuilt-invoice"), [])
        # Wrong document type -> no rows.
        self.assertEqual(database.get_filtered("2026-07-24", "2026-07-24", "prebuilt-receipt"), [])

    def test_get_all_parses_json_columns(self):
        database.save("t.pdf", "prebuilt-invoice", "p", "text", '{"A": 1}',
                      None, "ans", 1, 2, 3, "2026-07-23 10:00:00", "pid")
        rows = database.get_all()
        self.assertEqual(rows[0]["fields"], {"A": 1})
        self.assertIsNone(rows[0]["aiAnswerJson"])
        self.assertEqual(rows[0]["answer"], "ans")
        self.assertEqual(rows[0]["promptId"], "pid")

    def test_ensure_table_is_idempotent(self):
        database.ensure_table()  # second run must not raise
        self.assertIn("SQLite", database.describe())


if __name__ == "__main__":
    unittest.main()
